import concurrent.futures
import csv
import os
import sys
import tempfile
import urllib.request
import xml.etree.ElementTree as ET
from datetime import datetime
from urllib.parse import urlparse

import openpyxl
from openpyxl.utils import get_column_letter

try:
    from PySide6.QtCore import QThread, Qt, Signal
    from PySide6.QtGui import QGuiApplication
    from PySide6.QtWidgets import (
        QApplication,
        QCheckBox,
        QComboBox,
        QFileDialog,
        QGridLayout,
        QGroupBox,
        QHBoxLayout,
        QLabel,
        QLineEdit,
        QMainWindow,
        QMessageBox,
        QPlainTextEdit,
        QProgressBar,
        QPushButton,
        QScrollArea,
        QSizePolicy,
        QVBoxLayout,
        QWidget,
    )
except ImportError as error:
    raise ImportError(
        "Brak PySide6. Zainstaluj: pip install PySide6 openpyxl"
    ) from error

# >>> WBUDOWANY-BLOK-START: korekta-znakow
# Kod poniżej jest WKLEJONY automatycznie z desc_cleaner.py przez
# gen_skrypty_opisow.py. Nie edytuj go tutaj — zmiany nanoś w desc_cleaner.py
# i uruchom generator ponownie. Skrypt jest samowystarczalny: nie potrzebuje
# żadnych plików obok siebie.
import html
import re
from dataclasses import dataclass
from html.parser import HTMLParser
from typing import List, Optional, Tuple

# ---------------------------------------------------------------------------
# 1. KOREKTA ZNAKOW  (z "tlumaczenia v2.py")
# ---------------------------------------------------------------------------

EMOJI_PATTERN = re.compile(
    "["
    "\U0001F600-\U0001F64F"
    "\U0001F300-\U0001F5FF"
    "\U0001F680-\U0001F6FF"
    "\U0001F700-\U0001F77F"
    "\U0001F780-\U0001F7FF"
    "\U0001F800-\U0001F8FF"
    "\U0001F900-\U0001F9FF"
    "\U0001FA00-\U0001FA6F"
    "\U0001FA70-\U0001FAFF"
    "\U00002702-\U000027B0"
    "\U000024C2-\U0001F251"
    "\U0001f926-\U0001f937"
    "\U00010000-\U0010ffff"
    "♀-♂"
    "☀-⭕"
    "‍"
    "⏏"
    "⏩"
    "⌚"
    "️"
    "〰"
    "]+",
    flags=re.UNICODE,
)

POLISH_CHAR_MAP = {
    '&#378;ó&#322;ty': 'żółty',

    '&Aacute;': 'Ą', '&Cacute;': 'Ć', '&Eacute;': 'Ę',
    '&Lacute;': 'Ł', '&Nacute;': 'Ń', '&Oacute;': 'Ó',
    '&Sacute;': 'Ś', '&Zacute;': 'Ź', '&Zdot;': 'Ż',

    '&aacute;': 'ą', '&cacute;': 'ć', '&eacute;': 'ę',
    '&lacute;': 'ł', '&nacute;': 'ń', '&oacute;': 'ó',
    '&sacute;': 'ś', '&zacute;': 'ź', '&zdot;': 'ż',

    '&#260;': 'Ą', '&#262;': 'Ć', '&#280;': 'Ę', '&#321;': 'Ł',
    '&#323;': 'Ń', '&#211;': 'Ó', '&#346;': 'Ś', '&#377;': 'Ź',
    '&#379;': 'Ż',

    '&#261;': 'ą', '&#263;': 'ć', '&#281;': 'ę', '&#322;': 'ł',
    '&#324;': 'ń', '&#243;': 'ó', '&#347;': 'ś', '&#378;': 'ź',
    '&#380;': 'ż',

    '&deg;': '°', '&bull;': '•', '&ndash;': '–', '&rsquo;': '’',
    '&bdquo;': '„', '&rdquo;': '”',
    '&#10036;&#65039;': '', '&#10035;&#65039;': '', '&#9851;&#65039;': '',
    '&#128209;': '', '&#8222;': '„', '&#8221;': '”',
    '&#8216;': '‘', '&#8217;': '’', '&#8211;': '–', '&#8203;': '',
    '&#9989;': '', '&#9749;': '', '&#11088;': '', '&#10003;': '',
    '&#34;': '"', '&#39;': "'", '&#x2013;': '–', '&#2013;': '–',
    '&#2019;': '’', '&nbsp;': ' ', '&#178;': '²',
    '&#8220;': '“', '&#8230;': '…', '&#9679;': '•',

    '✔': '', '✅': '', '❓': '', '▶️': '',
    '⭐': '', '⚡': '', '➡': '',
}
# Najdluzsze klucze najpierw - zeby "&#10036;&#65039;" poszlo przed "&#10036;".
POLISH_CHAR_MAP_SORTED = dict(
    sorted(POLISH_CHAR_MAP.items(), key=lambda item: len(item[0]), reverse=True)
)

# Encje strukturalne HTML - rozwijane tylko wtedy, gdy tekst nie idzie juz
# przez parser HTML (inaczej zakodowany tekst stalby sie prawdziwym tagiem).
_MARKUP_ENTITIES = (
    ('&amp;', '&'), ('&lt;', '<'), ('&gt;', '>'),
    ('&quot;', '"'), ('&apos;', "'"),
)


def correct_text(text, keep_markup: bool = False):
    """
    Poprawia znaki: encje HTML -> polskie litery, usuwa emoji i smieci.

    keep_markup=True zostawia &lt; &gt; &amp; nietkniete.
    """
    if not isinstance(text, str) or not text:
        return text

    corrected = text
    for wrong, good in POLISH_CHAR_MAP_SORTED.items():
        if wrong in corrected:
            corrected = corrected.replace(wrong, good)

    if not keep_markup:
        for wrong, good in _MARKUP_ENTITIES:
            corrected = corrected.replace(wrong, good)

    corrected = EMOJI_PATTERN.sub('', corrected)
    return corrected.lstrip()
# <<< WBUDOWANY-BLOK-KONIEC: korekta-znakow

MAX_WORKERS = 10
DEFAULT_OUTPUT_DIR = os.path.join(os.path.expanduser("~"), "Downloads")
BASE_FIELDS = [
    "id",
    "id_bl",
    "url",
    "price",
    "avail",
    "weight",
    "stock",
    "cat",
    "name",
    "desc",
]
# W feedach id_bl siedzi jako <attrs><a name="id_bl">..., wyciagamy go na stala kolumne.
ID_BL_ATTR = "id_bl"
CSV_DELIMITER = "|"
FILTER_MODE_INCLUDE = "include"
FILTER_MODE_EXCLUDE = "exclude"


def clean_text(text):
    """Zastępuje znaki nowej linii i inne białe znaki pojedynczą spacją."""
    if not text:
        return ""
    return " ".join(text.split())


def clean_field(text):
    """
    clean_text + korekta znaków — dla pól tekstowych (cat, name, desc, atrybuty).

    URL-i, cen i identyfikatorów nie ruszamy: podmiana encji mogłaby uszkodzić
    adres, a emoji tam nie występują.
    """
    return correct_text(clean_text(text)) or ""


def load_filter_ids(file_path):
    """
    Wczytuje listę ID z pliku CSV/tekstowego (jedna wartość w wierszu, brana z
    pierwszej kolumny). Zwraca (zbiór_id, None) lub (set(), powod_bledu).
    """
    try:
        ids = set()
        with open(file_path, "r", encoding="utf-8-sig", newline="") as handle:
            for line in handle:
                token = line.strip()
                if not token:
                    continue
                for delimiter in (",", ";", "\t", "|"):
                    if delimiter in token:
                        token = token.split(delimiter, 1)[0].strip()
                        break
                if token:
                    ids.add(token)
        # Usuń typowy nagłówek kolumny, jeśli się pojawił.
        ids.discard("id")
        ids.discard("ID")
        ids.discard("Id")
        return ids, None
    except Exception as error:
        return set(), str(error)


def is_available(row):
    """Sprawdza, czy wiersz ma avail = 1 (produkt dostępny od ręki)."""
    return (row.get("avail") or "").strip() == "1"


def recalculate_columns(rows):
    """
    Przelicza zbiór atrybutów i maksymalną liczbę obrazów na podstawie
    zachowanych wierszy, aby uniknąć pustych kolumn po odfiltrowaniu.
    """
    attributes = set()
    max_images = 0
    for row in rows:
        for key in row:
            if key in BASE_FIELDS:
                continue
            if key.startswith("image") and key[5:].isdigit():
                max_images = max(max_images, int(key[5:]) + 1)
            else:
                attributes.add(key)
    return attributes, max_images


def download_xml(url, target_path):
    """Pobiera plik XML z podanego URL. Zwraca (True, None) lub (False, powód)."""
    try:
        opener = urllib.request.build_opener()
        opener.addheaders = [("User-agent", "Mozilla/5.0")]
        urllib.request.install_opener(opener)
        urllib.request.urlretrieve(url, target_path)
        return True, None
    except Exception as error:
        return False, str(error)


def parse_xml(file_path):
    """
    Parsuje plik XML i ekstrahuje dane produktowe.
    Zwraca (atrybuty, maks_obrazow, dane, None) lub ([], 0, [], powod_bledu).

    Pola tekstowe przechodzą korektę znaków (encje HTML -> polskie litery,
    usuwanie emoji) — tak samo jak w "tlumaczenia v2.py".
    """
    try:
        tree = ET.parse(file_path)
        root = tree.getroot()
        attributes = set()
        max_images = 0
        rows = []

        for element in root.findall("o"):
            cat_elem = element.find("cat")
            name_elem = element.find("name")
            desc_elem = element.find("desc")

            row = {
                "id": element.get("id"),
                "id_bl": "",
                "url": element.get("url"),
                "price": element.get("price"),
                "avail": element.get("avail"),
                "weight": element.get("weight"),
                "stock": element.get("stock"),
                "cat": clean_field(cat_elem.text) if cat_elem is not None else "",
                "name": clean_field(name_elem.text) if name_elem is not None else "",
                "desc": clean_field(desc_elem.text) if desc_elem is not None else "",
            }

            attrs_elem = element.find("attrs")
            if attrs_elem is not None:
                for attr in attrs_elem.findall("a"):
                    attr_name = attr.get("name")
                    if not attr_name:
                        continue
                    if attr_name == ID_BL_ATTR:
                        row["id_bl"] = clean_text(attr.text)
                        continue
                    if attr_name in BASE_FIELDS:
                        continue
                    attributes.add(attr_name)
                    row[attr_name] = clean_field(attr.text)

            images_in_row = 0
            imgs_elem = element.find("imgs")
            if imgs_elem is not None:
                main_image = imgs_elem.find("main")
                if main_image is not None and main_image.get("url"):
                    row["image0"] = main_image.get("url")
                    images_in_row = 1

                start_index = 1 if "image0" in row else 0
                for i, img in enumerate(imgs_elem.findall("i"), start=start_index):
                    if img.get("url"):
                        row[f"image{i}"] = img.get("url")
                        images_in_row = max(images_in_row, i + 1)

            max_images = max(max_images, images_in_row)
            rows.append(row)

        return sorted(attributes), max_images, rows, None

    except FileNotFoundError as error:
        return [], 0, [], f"Nie znaleziono pliku: {file_path} ({error})"
    except ET.ParseError as error:
        return [], 0, [], f"Błąd parsowania XML w {os.path.basename(file_path)}: {error}"
    except Exception as error:
        return [], 0, [], f"Nieoczekiwany błąd parsowania: {error}"


def write_csv(rows, attributes, max_images, file_path):
    """Zapisuje dane jednego feedu do pliku CSV."""
    fields = BASE_FIELDS + list(attributes) + [f"image{i}" for i in range(max_images)]
    try:
        with open(file_path, "w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(
                handle, fieldnames=fields, delimiter=CSV_DELIMITER, extrasaction="ignore"
            )
            writer.writeheader()
            writer.writerows(rows)
        return True, None
    except Exception as error:
        return False, str(error)


def save_error_report(download_errors, parse_errors, write_errors, output_dir):
    """Zapisuje raport błędów (pobieranie/parsowanie/zapis) do pliku XLSX. Zwraca ścieżkę lub None."""
    if not download_errors and not parse_errors and not write_errors:
        return None

    workbook = openpyxl.Workbook()
    sheets = []

    if download_errors:
        sheet = workbook.active
        sheet.title = "Bledy pobierania"
        sheet.append(["Nieudany URL", "Powód błędu", "Plik"])
        for url, reason, file_name in download_errors:
            sheet.append([url, reason, file_name])
        sheets.append(sheet)
    else:
        workbook.remove(workbook.active)

    if parse_errors:
        sheet = workbook.create_sheet("Bledy parsowania")
        sheet.append(["URL", "Powód błędu", "Plik"])
        for url, reason, file_name in parse_errors:
            sheet.append([url, reason, file_name])
        sheets.append(sheet)

    if write_errors:
        sheet = workbook.create_sheet("Bledy zapisu")
        sheet.append(["Plik CSV", "Powód błędu"])
        for file_name, reason in write_errors:
            sheet.append([file_name, reason])
        sheets.append(sheet)

    for sheet in sheets:
        for column in sheet.columns:
            max_length = 0
            letter = get_column_letter(column[0].column)
            for cell in column:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(value))
            sheet.column_dimensions[letter].width = min(max_length + 2, 80)

    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    report_path = os.path.join(output_dir, f"RAPORT_BLEDOW_XMLCSV_{timestamp}.xlsx")
    workbook.save(report_path)
    return report_path


def download_and_parse_url(url):
    """Pobiera i parsuje jeden URL. Przeznaczone do uruchamiania w osobnym wątku."""
    temp_dir = tempfile.gettempdir()
    file_name = os.path.basename(urlparse(url).path) or f"feed_{abs(hash(url))}.xml"
    base_name = os.path.splitext(file_name)[0]
    temp_stamp = datetime.now().strftime("%Y%m%d%H%M%S%f")
    local_xml_path = os.path.join(temp_dir, f"temp_{base_name}_{temp_stamp}.xml")

    success, error_message = download_xml(url, local_xml_path)
    if not success:
        return "download_error", (url, error_message, file_name)

    attributes, max_images, rows, parse_error = parse_xml(local_xml_path)

    try:
        os.remove(local_xml_path)
    except OSError:
        pass

    if parse_error:
        return "parse_error", (url, parse_error, file_name)

    if not rows:
        return "parse_error", (url, "Brak elementów <o> po parsowaniu", file_name)

    return "success", (rows, attributes, max_images, base_name, file_name)


class ProcessorThread(QThread):
    progress_signal = Signal(str, float, str)
    done_signal = Signal(dict)
    error_signal = Signal(str)

    def __init__(
        self,
        urls,
        output_dir,
        filter_ids=None,
        avail_only=False,
        filter_mode=FILTER_MODE_INCLUDE,
    ):
        super().__init__()
        self.urls = urls
        self.output_dir = output_dir
        self.filter_ids = set(filter_ids) if filter_ids else set()
        self.avail_only = bool(avail_only)
        self.filter_mode = filter_mode

    def _apply_filters(self, rows):
        """
        Filtruje wiersze jednego feedu (avail, potem lista ID).
        Zwraca (wiersze, liczba_przed, liczba_po_avail).
        """
        rows_before = len(rows)
        if self.avail_only:
            rows = [row for row in rows if is_available(row)]
        rows_after_avail = len(rows)

        if self.filter_ids:
            if self.filter_mode == FILTER_MODE_EXCLUDE:
                rows = [row for row in rows if row.get("id") not in self.filter_ids]
            else:
                rows = [row for row in rows if row.get("id") in self.filter_ids]

        return rows, rows_before, rows_after_avail

    def _emit_progress(self, message, value, tone="normal"):
        self.progress_signal.emit(message, value, tone)

    def run(self):
        try:
            if not self.urls:
                self.error_signal.emit("Podaj co najmniej jeden URL pliku XML.")
                return

            if not os.path.exists(self.output_dir):
                try:
                    os.makedirs(self.output_dir, exist_ok=True)
                except Exception as error:
                    self.error_signal.emit(
                        f"Nie można utworzyć katalogu zapisu:\n{self.output_dir}\n{error}"
                    )
                    return

            saved = 0
            download_errors = []
            parse_errors = []
            write_errors = []
            filtered_out = []
            rows_written = 0
            rows_before_filter = 0
            rows_after_avail_total = 0
            total = len(self.urls)

            filtering_active = self.avail_only or bool(self.filter_ids)
            exclude_mode = self.filter_mode == FILTER_MODE_EXCLUDE

            self._emit_progress(
                f"Rozpoczynam przetwarzanie {total} linków (max {MAX_WORKERS} wątków)...", 0.02
            )

            with concurrent.futures.ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
                futures = [executor.submit(download_and_parse_url, url) for url in self.urls]
                for index, future in enumerate(concurrent.futures.as_completed(futures), start=1):
                    progress = 0.02 + (0.96 * index / total)
                    try:
                        status, data = future.result()
                        if status == "success":
                            rows, attributes, max_images, base_name, _ = data

                            rows, feed_before, feed_after_avail = self._apply_filters(rows)
                            rows_before_filter += feed_before
                            rows_after_avail_total += feed_after_avail
                            if filtering_active:
                                recalculated, max_images = recalculate_columns(rows)
                                attributes = sorted(recalculated)

                            if not rows:
                                filtered_out.append((base_name, feed_before))
                                self._emit_progress(
                                    f"Filtry odrzuciły wszystkie {feed_before} wierszy "
                                    f"{index}/{total}: {base_name} — pomijam",
                                    progress,
                                    "warn",
                                )
                                continue

                            timestamp = datetime.now().strftime("%d%m%y-%H%M%S")
                            csv_path = os.path.join(
                                self.output_dir, f"{base_name}_{timestamp}.csv"
                            )
                            self._emit_progress(
                                f"Zapisywanie {index}/{total}: {os.path.basename(csv_path)}",
                                progress,
                            )
                            save_ok, save_error = write_csv(rows, attributes, max_images, csv_path)
                            if save_ok:
                                saved += 1
                                rows_written += len(rows)
                            else:
                                write_errors.append((csv_path, save_error))
                                self._emit_progress(
                                    f"Błąd zapisu {index}/{total}: {os.path.basename(csv_path)}",
                                    progress,
                                    "warn",
                                )
                        elif status == "download_error":
                            download_errors.append(data)
                            self._emit_progress(
                                f"Błąd pobierania {index}/{total}: {data[2]}", progress, "warn"
                            )
                        else:
                            parse_errors.append(data)
                            self._emit_progress(
                                f"Błąd parsowania {index}/{total}: {data[2]}", progress, "warn"
                            )
                    except Exception as error:
                        parse_errors.append(("?", f"Błąd krytyczny wątku: {error}", "?"))
                        self._emit_progress(f"Błąd krytyczny wątku {index}/{total}", progress, "warn")

            report_path = save_error_report(
                download_errors, parse_errors, write_errors, self.output_dir
            )
            error_count = len(download_errors) + len(parse_errors) + len(write_errors)

            summary_lines = [
                f"Zapisane pliki CSV: {saved}/{total}",
                f"Wiersze produktów: {rows_written}",
            ]

            if self.avail_only:
                summary_lines.append(
                    f"Filtr avail=1: {rows_after_avail_total}/{rows_before_filter} wierszy"
                )
            if self.filter_ids:
                summary_lines.append(
                    f"Filtr ID ({'wyklucz' if exclude_mode else 'dołącz'}): "
                    f"{rows_written}/{rows_after_avail_total} wierszy "
                    f"(lista: {len(self.filter_ids)} ID)"
                )
            if filtered_out:
                summary_lines.append(
                    f"Pliki pominięte (0 wierszy po filtrach): {len(filtered_out)}"
                )

            summary_lines.extend(
                [
                    f"Błędy pobierania: {len(download_errors)}",
                    f"Błędy parsowania: {len(parse_errors)}",
                    f"Błędy zapisu: {len(write_errors)}",
                ]
            )

            if saved > 0:
                summary_lines.append(f"\nFolder zapisu:\n{os.path.abspath(self.output_dir)}")
            if report_path:
                summary_lines.append(f"Raport błędów: {os.path.basename(report_path)}")

            ok = saved > 0 and error_count == 0
            if saved > 0:
                title = "Sukces" if error_count == 0 else "Zakończono z błędami"
                tone = "ok" if error_count == 0 else "warn"
            elif filtered_out and error_count == 0:
                title = "Brak danych po filtrach"
                tone = "warn"
                summary_lines.append(
                    f"\nWszystkie {rows_before_filter} wierszy zostało odfiltrowanych.\n"
                    f"Filtr avail=1: {'tak' if self.avail_only else 'nie'}\n"
                    f"Filtr ID: "
                    f"{('wyklucz' if exclude_mode else 'dołącz') if self.filter_ids else 'nie'}"
                )
            else:
                title = "Brak danych"
                tone = "warn"

            self.done_signal.emit(
                {
                    "ok": ok,
                    "title": title,
                    "message": "\n".join(summary_lines),
                    "tone": tone,
                    "progress": 1.0,
                }
            )
        except Exception as error:
            self.error_signal.emit(str(error))


# Paleta i styl przeniesione z google_sheets_translate_excel_gui_multilang.py
# (jasne tlo / rozowe akcenty) — wszystkie narzedzia maja wygladac tak samo.
ACCENT = "#ff69b4"
ACCENT_HOVER = "#e754a7"
APP_BG = "#fff7fb"
PANEL_BG = "#ffffff"
INPUT_BG = "#fff2f8"
BORDER = "#f7b3d2"
TEXT = "#3d2130"
MUTED = "#7d5a6b"
TROUGH = "#f5d7e6"
OK_COLOR = "#1f7a4c"
WARN_COLOR = "#a35300"

QSS = f"""
QWidget {{ background-color: {APP_BG}; color: {TEXT}; font-size: 13px; }}
QGroupBox {{
    background-color: {PANEL_BG};
    border: 1px solid {BORDER};
    border-radius: 12px;
    margin-top: 16px;
    padding: 10px 10px 8px 10px;
    font-weight: bold;
}}
QGroupBox::title {{
    subcontrol-origin: margin;
    subcontrol-position: top left;
    left: 12px;
    padding: 2px 6px;
    color: {ACCENT_HOVER};
}}
QLabel {{ background: transparent; color: {TEXT}; font-weight: normal; }}
QLabel#Hint {{ color: {MUTED}; }}
QLineEdit, QPlainTextEdit, QComboBox {{
    background-color: {INPUT_BG};
    border: 1px solid {BORDER};
    border-radius: 8px;
    padding: 5px 8px;
    color: {TEXT};
    selection-background-color: {ACCENT};
    selection-color: #ffffff;
}}
QLineEdit:focus, QPlainTextEdit:focus {{ border: 1px solid {ACCENT}; }}
QLineEdit:disabled {{ background-color: {TROUGH}; color: {MUTED}; }}
QComboBox::drop-down {{ border: none; width: 24px; }}
QComboBox QAbstractItemView {{
    background: {PANEL_BG};
    color: {TEXT};
    border: 1px solid {BORDER};
    selection-background-color: {ACCENT};
    selection-color: #ffffff;
    outline: none;
}}
QPushButton {{
    background-color: {ACCENT};
    color: #ffffff;
    border: none;
    border-radius: 10px;
    padding: 9px 16px;
    font-weight: bold;
}}
QPushButton:hover {{ background-color: {ACCENT_HOVER}; }}
QPushButton:pressed {{ background-color: {ACCENT_HOVER}; }}
QPushButton:disabled {{ background-color: {TROUGH}; color: #ffffff; }}
QCheckBox {{ background: transparent; color: {TEXT}; spacing: 8px; font-weight: normal; }}
QCheckBox::indicator {{
    width: 16px; height: 16px;
    border: 1px solid {BORDER}; border-radius: 4px; background: {INPUT_BG};
}}
QCheckBox::indicator:checked {{ background: {ACCENT}; border: 1px solid {ACCENT}; }}
QProgressBar {{
    background-color: {TROUGH};
    border: 1px solid {BORDER};
    border-radius: 8px;
    text-align: center;
    color: {TEXT};
    min-height: 18px;
}}
QProgressBar::chunk {{ background-color: {ACCENT}; border-radius: 7px; }}
QScrollArea {{ background: transparent; border: none; }}
QScrollArea > QWidget > QWidget {{ background: transparent; }}
QScrollBar:vertical {{ background: {INPUT_BG}; width: 12px; margin: 0; }}
QScrollBar::handle:vertical {{ background: {BORDER}; border-radius: 6px; min-height: 24px; }}
QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical {{ height: 0; }}
QScrollBar:horizontal {{ background: {INPUT_BG}; height: 12px; margin: 0; }}
QScrollBar::handle:horizontal {{ background: {BORDER}; border-radius: 6px; min-width: 24px; }}
QScrollBar::add-line:horizontal, QScrollBar::sub-line:horizontal {{ width: 0; }}
QScrollBar::add-page, QScrollBar::sub-page {{ background: transparent; }}
"""

WINDOW_TITLE = "Konwerter XML → CSV (osobne pliki)"
WINDOW_SUBTITLE = "Pobierz wiele feedów XML jednocześnie i zapisz każdy do osobnego pliku CSV."
RUN_BUTTON_TEXT = "Przetwórz na OSOBNE pliki CSV"
OUTPUT_LABEL = "Folder zapisu CSV:"
FILTER_PLACEHOLDER = "Brak filtra – w CSV znajdą się wszystkie wiersze"
FILTER_MODE_TOOLTIP = "Wskazane ID trafiają do CSV (dołącz) albo są z niego usuwane (wyklucz)."
AVAIL_TOOLTIP = "Do CSV trafią wyłącznie wiersze, w których atrybut avail ma wartość 1."


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(WINDOW_TITLE)

        self.worker = None
        self.filter_ids = set()
        self.filter_path = None

        # Cala zawartosc siedzi w jednym scrollu strony — tak samo jak w
        # google_sheets_translate_excel_gui_multilang.py, zeby na niskich
        # ekranach zadna sekcja nie byla nieosiagalna.
        content = QWidget()
        root = QVBoxLayout(content)

        hint = QLabel(WINDOW_SUBTITLE)
        hint.setObjectName("Hint")
        hint.setWordWrap(True)
        root.addWidget(hint)

        root.addWidget(self._build_urls_group(), stretch=1)
        root.addWidget(self._build_settings_group())
        root.addLayout(self._build_action_row())
        root.addWidget(self._build_progress_group())

        page = QScrollArea()
        page.setWidgetResizable(True)
        page.setFrameShape(QScrollArea.NoFrame)
        page.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        page.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        page.setWidget(content)
        self.setCentralWidget(page)

        # Okno nigdy wieksze niz dostepny obszar ekranu (bez paska zadan).
        screen = QGuiApplication.primaryScreen()
        avail = screen.availableGeometry() if screen else None
        if avail is not None:
            self.resize(min(880, avail.width() - 40), min(780, avail.height() - 60))
            self.move(
                avail.left() + max(0, (avail.width() - self.width()) // 2),
                avail.top() + max(0, (avail.height() - self.height()) // 2),
            )
        else:
            self.resize(880, 780)

    # ---------- budowa UI ----------
    def _build_urls_group(self) -> QGroupBox:
        g = QGroupBox("Linki XML")
        lay = QVBoxLayout(g)
        lay.addWidget(QLabel("Wklej URL-e plików XML, każdy w nowej linii."))

        self.url_input = QPlainTextEdit()
        self.url_input.setPlaceholderText(
            "https://example.com/feed.xml\nhttps://example.com/feed2.xml"
        )
        self.url_input.setMinimumHeight(160)
        self.url_input.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        lay.addWidget(self.url_input, 1)
        return g

    def _build_settings_group(self) -> QGroupBox:
        g = QGroupBox("Ustawienia")
        grid = QGridLayout(g)

        grid.addWidget(QLabel(OUTPUT_LABEL), 0, 0)
        self.output_input = QLineEdit(DEFAULT_OUTPUT_DIR)
        grid.addWidget(self.output_input, 0, 1)
        self.output_btn = QPushButton("Wybierz folder")
        self.output_btn.clicked.connect(self.pick_output_dir)
        grid.addWidget(self.output_btn, 0, 2)

        grid.addWidget(QLabel("Filtr ID (CSV, opcjonalnie):"), 1, 0)
        self.filter_input = QLineEdit()
        self.filter_input.setReadOnly(True)
        self.filter_input.setPlaceholderText(FILTER_PLACEHOLDER)
        grid.addWidget(self.filter_input, 1, 1)

        self.filter_btn = QPushButton("Wybierz plik CSV")
        self.filter_btn.clicked.connect(self.pick_filter_file)
        self.filter_clear_btn = QPushButton("Wyczyść")
        self.filter_clear_btn.clicked.connect(self.clear_filter_file)

        filter_buttons = QWidget()
        filter_buttons_layout = QHBoxLayout(filter_buttons)
        filter_buttons_layout.setContentsMargins(0, 0, 0, 0)
        filter_buttons_layout.setSpacing(8)
        filter_buttons_layout.addWidget(self.filter_btn)
        filter_buttons_layout.addWidget(self.filter_clear_btn)
        grid.addWidget(filter_buttons, 1, 2)

        grid.addWidget(QLabel("Tryb filtra ID:"), 2, 0)
        self.filter_mode_combo = QComboBox()
        self.filter_mode_combo.addItem("Zostaw tylko wskazane ID", FILTER_MODE_INCLUDE)
        self.filter_mode_combo.addItem("Wyklucz wskazane ID", FILTER_MODE_EXCLUDE)
        self.filter_mode_combo.setToolTip(FILTER_MODE_TOOLTIP)
        grid.addWidget(self.filter_mode_combo, 2, 1, 1, 2)

        self.avail_only_check = QCheckBox("Tylko dostępne produkty (avail = 1)")
        self.avail_only_check.setToolTip(AVAIL_TOOLTIP)
        grid.addWidget(self.avail_only_check, 3, 1, 1, 2)

        grid.setColumnStretch(1, 1)
        return g

    def _build_action_row(self) -> QHBoxLayout:
        row = QHBoxLayout()
        self.run_btn = QPushButton(RUN_BUTTON_TEXT)
        self.run_btn.clicked.connect(self.run_processing)
        row.addWidget(self.run_btn)
        return row

    def _build_progress_group(self) -> QGroupBox:
        g = QGroupBox("Postęp")
        lay = QVBoxLayout(g)

        self.progress = QProgressBar()
        self.progress.setRange(0, 1000)
        self.progress.setValue(0)
        lay.addWidget(self.progress)

        self.status = QLabel("Gotowy.")
        self.status.setObjectName("Status")
        self.status.setWordWrap(True)
        lay.addWidget(self.status)
        return g

    # ---------- akcje ----------
    def pick_output_dir(self):
        selected = QFileDialog.getExistingDirectory(
            self, "Wybierz folder zapisu", self.output_input.text().strip() or DEFAULT_OUTPUT_DIR
        )
        if selected:
            self.output_input.setText(selected)

    def pick_filter_file(self):
        start_dir = self.output_input.text().strip() or DEFAULT_OUTPUT_DIR
        selected, _ = QFileDialog.getOpenFileName(
            self,
            "Wybierz plik CSV z listą ID",
            start_dir,
            "Pliki CSV (*.csv);;Pliki tekstowe (*.txt);;Wszystkie pliki (*)",
        )
        if not selected:
            return

        ids, error = load_filter_ids(selected)
        if error:
            QMessageBox.warning(self, "Błąd", f"Nie można wczytać pliku:\n{error}")
            return
        if not ids:
            QMessageBox.warning(self, "Uwaga", "Wybrany plik nie zawiera żadnych ID.")
            return

        self.filter_path = selected
        self.filter_ids = ids
        self.filter_input.setText(f"{os.path.basename(selected)} — {len(ids)} ID")

    def clear_filter_file(self):
        self.filter_path = None
        self.filter_ids = set()
        self.filter_input.clear()

    def run_processing(self):
        urls = [line.strip() for line in self.url_input.toPlainText().splitlines() if line.strip()]
        output_dir = self.output_input.text().strip() or DEFAULT_OUTPUT_DIR

        if not urls:
            QMessageBox.warning(self, "Błąd", "Musisz podać co najmniej jeden URL pliku XML.")
            return

        self.output_input.setText(output_dir)
        self.run_btn.setEnabled(False)
        self.run_btn.setText("Przetwarzanie...")
        self.status.setText("Start przetwarzania...")
        self.progress.setValue(10)

        self.worker = ProcessorThread(
            urls,
            output_dir,
            self.filter_ids,
            self.avail_only_check.isChecked(),
            self.filter_mode_combo.currentData(),
        )
        self.worker.progress_signal.connect(self.on_progress)
        self.worker.done_signal.connect(self.on_done)
        self.worker.error_signal.connect(self.on_error)
        self.worker.start()

    def on_progress(self, message, value, tone):
        color = WARN_COLOR if tone == "warn" else MUTED
        self.status.setStyleSheet(f"color: {color}; font-weight: 600;")
        self.status.setText(message)
        self.progress.setValue(max(0, min(1000, int(value * 1000))))

    def on_done(self, payload):
        tone = payload.get("tone", "ok")
        color = OK_COLOR if tone == "ok" else WARN_COLOR
        self.status.setStyleSheet(f"color: {color}; font-weight: 700;")
        self.status.setText(payload.get("title", "Zakończono"))
        self.progress.setValue(int(payload.get("progress", 1.0) * 1000))

        if payload.get("ok", False):
            QMessageBox.information(self, payload.get("title", "Sukces"), payload.get("message", ""))
        else:
            QMessageBox.warning(self, payload.get("title", "Uwaga"), payload.get("message", ""))

        self._reset_run_button()

    def on_error(self, message):
        self.status.setStyleSheet(f"color: {WARN_COLOR}; font-weight: 700;")
        self.status.setText("Błąd krytyczny")
        self.progress.setValue(0)
        self._reset_run_button()
        QMessageBox.critical(self, "Błąd", message)

    def _reset_run_button(self):
        self.run_btn.setEnabled(True)
        self.run_btn.setText(RUN_BUTTON_TEXT)


def main():
    os.environ.setdefault("QT_SCALE_FACTOR", "0.9")

    app = QApplication(sys.argv)
    app.setStyleSheet(QSS)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
