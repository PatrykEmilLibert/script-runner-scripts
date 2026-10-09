from __future__ import annotations

import json
import math
import threading
import time
import base64
import re
from urllib.parse import urlencode
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from dataclasses import dataclass
from collections import deque
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Any

import customtkinter as ctk
import pandas as pd
import requests
from bol_api_config_store import BolApiConfigStore, BolApiSettings
from openpyxl import load_workbook
from tkinter import filedialog, messagebox


APP_TITLE = "BOL Offers Creator (MVP)"
DEFAULT_API_URL = "https://api.bol.com/retailer/offers"
DEFAULT_UPDATE_API_URL = "https://api.bol.com/retailer/offers/{offer-id}"
DEFAULT_AUTH_URL = "https://login.bol.com/token"
DEFAULT_ECONOMIC_OPERATORS_API_URL = "https://api.bol.com/retailer/economic-operators"
SMPRODS_BASE_URL = "https://api-sm-prods.sm-prods.com"
SMPRODS_API_TOKEN = "y7SeKeGSfVZtH9dCxwVULWTbcfWBrVq2WKcJssq8Pz8o5t3DFDpQ12BGRGc1S3fOJ2UC3tRMi29ChrseLAsl4GhHKR3Y9ALr9Zfq8pyeYtlExRas7rOfvRTrqKdEOJ8y"
DEFAULT_LISTED_ENDPOINT = "supermerchant_bol_listed"
DEFAULT_LISTED_FILE_PATH = Path(str((Path(__file__).parent / "listed.csv").resolve()))
LISTED_PAGE_SIZE = 1000
LISTED_REQUEST_TIMEOUT = 300
LISTED_MAX_RETRIES = 4
DEFAULT_TEMPLATE_XLSX_PATH = Path(
    str((Path(__file__).parent / "furnitex_maxdywanik.xlsx").resolve())
)

ACCENT = "#ff69b4"
ACCENT_HOVER = "#e754a7"
APP_BG = "#fff7fb"
PANEL_BG = "#ffffff"
INPUT_BG = "#fff2f8"
BORDER = "#f7b3d2"
TEXT = "#3d2130"
MUTED = "#7d5a6b"
ERROR = "#b42318"
SUCCESS = "#2f7d4f"

SEND_RATE_LIMIT_PER_SEC = 40
SEND_MAX_WORKERS = 12
SEND_MAX_RETRIES = 4
SEND_REQUEST_TIMEOUT = 20
LOOKUP_RATE_LIMIT_PER_SEC = 10
LOOKUP_MAX_WORKERS = 10
LOOKUP_REQUEST_TIMEOUT = 20


@dataclass(slots=True)
class ColumnMapping:
    offer_id: str = "offer_id"
    ean: str = "EAN"
    price: str = "price"
    stock: str = "stock"
    reference: str = "id"
    on_hold: str = "on_hold"
    economic_operator_id: str = "economicOperatorId"


@dataclass(slots=True)
class OfferDefaults:
    condition_category: str = "NEW"
    fulfilment_method: str = "FBR"
    fulfilment_schedule: str = "BOL_DELIVERY_PROMISE"
    min_days_to_customer: int = 4
    max_days_to_customer: int = 8
    country_codes: tuple[str, str] = ("NL", "BE")
    managed_by_retailer: bool = True
    on_hold_by_retailer: bool = False


# Zakodowane na sztywno wg offers-v11 (dawniej wczytywane z offers-v11.yaml).
CREATE_REQUIRED_FIELDS: list[str] = ["condition", "pricing", "ean", "fulfilment"]
UPDATE_PATCH_FIELDS: list[str] = [
    "unknownProductTitle",
    "economicOperatorId",
    "onHoldByRetailer",
    "reference",
    "countryAvailabilities",
    "pricing",
    "fulfilment",
    "stock",
]


class RateLimiter:
    def __init__(self, rate_per_sec: int):
        self.rate_per_sec = rate_per_sec
        self.lock = threading.Lock()
        self.calls = deque()

    def acquire(self) -> None:
        while True:
            with self.lock:
                now = time.monotonic()
                while self.calls and now - self.calls[0] >= 1:
                    self.calls.popleft()

                if len(self.calls) < self.rate_per_sec:
                    self.calls.append(now)
                    return

                sleep_for = 1 - (now - self.calls[0])
            time.sleep(max(sleep_for, 0.001))


class OAuthTokenManager:
    def __init__(self, auth_url: str, client_id: str, client_secret: str):
        self.auth_url = auth_url
        self.client_id = client_id
        self.client_secret = client_secret
        self.lock = threading.Lock()
        self.token: str | None = None
        self.expires_in = 0
        self.token_time = 0.0

    def get_token(self) -> str:
        with self.lock:
            now = time.time()
            if self.token and now < self.token_time + self.expires_in - 30:
                return self.token

            auth_raw = f"{self.client_id}:{self.client_secret}".encode("utf-8")
            auth_encoded = base64.b64encode(auth_raw).decode("utf-8")

            headers = {
                "Content-Type": "application/x-www-form-urlencoded",
                "Accept": "application/json",
                "Authorization": f"Basic {auth_encoded}",
            }

            response = requests.post(
                self.auth_url,
                params={"grant_type": "client_credentials"},
                headers=headers,
                timeout=15,
            )

            if response.status_code != 200:
                raise ValueError(f"Błąd pobierania tokena HTTP {response.status_code}: {response.text[:250]}")

            payload = response.json()
            access_token = payload.get("access_token")
            expires_in = int(payload.get("expires_in", 0))
            if not access_token:
                raise ValueError("Brak access_token w odpowiedzi OAuth")

            self.token = access_token
            self.expires_in = expires_in if expires_in > 0 else 300
            self.token_time = now
            return self.token


class OfferPayloadBuilder:
    @staticmethod
    def _round_price_2(value: float) -> float:
        return float(Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))

    @staticmethod
    def _normalize_ean(value: Any) -> str:
        if value is None:
            return ""

        if isinstance(value, float) and pd.isna(value):
            return ""

        raw = str(value).strip()
        if not raw or raw.lower() == "nan":
            return ""

        candidate = raw.replace(" ", "")
        digits = ""

        if candidate.isdigit():
            digits = candidate
        else:
            try:
                decimal_value = Decimal(candidate)
                if decimal_value != decimal_value.to_integral_value():
                    return ""
                digits = format(decimal_value.quantize(Decimal("1")), "f").split(".")[0]
            except (InvalidOperation, ValueError):
                return ""

        if not digits or len(digits) > 13:
            return ""

        return digits.zfill(13)

    @staticmethod
    def _has_value(value: Any) -> bool:
        if value is None:
            return False
        if isinstance(value, float) and pd.isna(value):
            return False
        value_str = str(value).strip()
        return bool(value_str and value_str.lower() != "nan")

    @staticmethod
    def _normalize_offer_id(value: Any) -> str:
        if not OfferPayloadBuilder._has_value(value):
            return ""
        return str(value).strip()

    @staticmethod
    def _to_bool(value: Any, fallback: bool) -> bool:
        if value is None:
            return fallback
        if isinstance(value, bool):
            return value
        lowered = str(value).strip().lower()
        if lowered in {"1", "true", "tak", "yes", "y"}:
            return True
        if lowered in {"0", "false", "nie", "no", "n"}:
            return False
        return fallback

    @staticmethod
    def _to_int(value: Any, fallback: int = 0) -> int:
        if value is None or str(value).strip() == "":
            return fallback
        return int(float(value))

    @staticmethod
    def _to_float(value: Any, fallback: float = 0.0) -> float:
        if value is None or str(value).strip() == "":
            return fallback
        if isinstance(value, str):
            value = value.replace(",", ".")
        return float(value)

    @staticmethod
    def _round_up_to_49_or_99(value: float) -> float:
        if value <= 0:
            return 0.0
        cents = int(math.ceil(value * 100))
        whole = cents // 100
        fractional = cents % 100

        if fractional <= 49:
            target_cents = whole * 100 + 49
        else:
            target_cents = whole * 100 + 99

        return target_cents / 100

    @staticmethod
    def build_create_payload(row: dict[str, Any], mapping: ColumnMapping, defaults: OfferDefaults) -> dict[str, Any]:
        ean_value = OfferPayloadBuilder._normalize_ean(row.get(mapping.ean, ""))
        if not ean_value:
            raise ValueError("Brak lub niepoprawny EAN (wymagane 13 cyfr)")

        condition_category = defaults.condition_category
        fulfilment_method = defaults.fulfilment_method
        price = OfferPayloadBuilder._to_float(row.get(mapping.price), fallback=0.0)
        if price <= 0:
            raise ValueError("Cena musi być > 0")

        main_price = OfferPayloadBuilder._round_up_to_49_or_99(price)
        bundle_prices = [
            {
                "quantity": 1,
                "unitPrice": main_price,
            }
        ]
        if main_price < 30:
            bundle_prices.extend(
                [
                    {
                        "quantity": 2,
                        "unitPrice": OfferPayloadBuilder._round_price_2(main_price * 0.98),
                    },
                    {
                        "quantity": 3,
                        "unitPrice": OfferPayloadBuilder._round_price_2(main_price * 0.97),
                    },
                    {
                        "quantity": 4,
                        "unitPrice": OfferPayloadBuilder._round_price_2(main_price * 0.96),
                    },
                ]
            )

        stock_amount = OfferPayloadBuilder._to_int(row.get(mapping.stock), fallback=0)
        on_hold = OfferPayloadBuilder._to_bool(row.get(mapping.on_hold), defaults.on_hold_by_retailer)

        fulfilment_payload: dict[str, Any] = {
            "method": fulfilment_method,
            "schedule": defaults.fulfilment_schedule,
        }
        if fulfilment_method == "FBR":
            fulfilment_payload["deliveryPromise"] = {
                "minimumDaysToCustomer": defaults.min_days_to_customer,
                "maximumDaysToCustomer": defaults.max_days_to_customer,
            }

        payload: dict[str, Any] = {
            "ean": ean_value,
            "condition": {"category": condition_category},
            "pricing": {
                "bundlePrices": bundle_prices
            },
            "fulfilment": fulfilment_payload,
            "countryAvailabilities": [{"countryCode": code} for code in defaults.country_codes],
            "onHoldByRetailer": on_hold,
        }

        reference = str(row.get(mapping.reference, "")).strip()
        if reference:
            payload["reference"] = reference

        economic_operator_id_raw = row.get(mapping.economic_operator_id)
        if OfferPayloadBuilder._has_value(economic_operator_id_raw):
            payload["economicOperatorId"] = str(economic_operator_id_raw).strip()

        payload["stock"] = {
            "amount": stock_amount,
            "managedByRetailer": defaults.managed_by_retailer,
        }

        return payload

    @staticmethod
    def build_update_request(
        row: dict[str, Any],
        mapping: ColumnMapping,
        defaults: OfferDefaults,
        require_offer_id: bool = True,
        multiplier: float = 1.0,
    ) -> tuple[str, dict[str, Any]]:
        offer_id = OfferPayloadBuilder._normalize_offer_id(row.get(mapping.offer_id, ""))
        if require_offer_id and not offer_id:
            raise ValueError("Brak offer-id w wierszu")

        payload: dict[str, Any] = {}

        reference_raw = row.get(mapping.reference)
        if OfferPayloadBuilder._has_value(reference_raw):
            payload["reference"] = str(reference_raw).strip()

        on_hold_raw = row.get(mapping.on_hold)
        if OfferPayloadBuilder._has_value(on_hold_raw):
            payload["onHoldByRetailer"] = OfferPayloadBuilder._to_bool(on_hold_raw, defaults.on_hold_by_retailer)

        economic_operator_id_raw = row.get(mapping.economic_operator_id)
        if OfferPayloadBuilder._has_value(economic_operator_id_raw):
            payload["economicOperatorId"] = str(economic_operator_id_raw).strip()

        stock_raw = row.get(mapping.stock)
        if OfferPayloadBuilder._has_value(stock_raw):
            payload["stock"] = {
                "amount": OfferPayloadBuilder._to_int(stock_raw, fallback=0),
                "managedByRetailer": defaults.managed_by_retailer,
            }

        price_raw = row.get(mapping.price)
        if OfferPayloadBuilder._has_value(price_raw):
            price = OfferPayloadBuilder._to_float(price_raw, fallback=0.0)
            if price <= 0:
                raise ValueError("Cena (update) musi być > 0, jeżeli podana")
            main_price = OfferPayloadBuilder._round_up_to_49_or_99(price * multiplier)
            bundle_prices = [{"quantity": 1, "unitPrice": main_price}]
            if main_price < 30:
                bundle_prices.extend(
                    [
                        {"quantity": 2, "unitPrice": OfferPayloadBuilder._round_price_2(main_price * 0.98)},
                        {"quantity": 3, "unitPrice": OfferPayloadBuilder._round_price_2(main_price * 0.97)},
                        {"quantity": 4, "unitPrice": OfferPayloadBuilder._round_price_2(main_price * 0.96)},
                    ]
                )
            payload["pricing"] = {"bundlePrices": bundle_prices}

        if not payload:
            raise ValueError("Brak pól do aktualizacji w wierszu (UPDATE wysyła tylko dane z pliku)")

        return offer_id, payload

    @staticmethod
    def build_payload(row: dict[str, Any], mapping: ColumnMapping, defaults: OfferDefaults) -> dict[str, Any]:
        return OfferPayloadBuilder.build_create_payload(row, mapping, defaults)


def _fetch_listed_multipliers(endpoint: str) -> dict[str, float]:
    url = f"{SMPRODS_BASE_URL}/marketplaces_paginated/{endpoint}"
    headers = {"Authorization": f"Bearer {SMPRODS_API_TOKEN}"}
    page = 1
    result: dict[str, float] = {}

    while True:
        payload: dict[str, Any] | None = None
        for attempt in range(1, LISTED_MAX_RETRIES + 1):
            try:
                r = requests.get(
                    f"{url}?page={page}&page_size={LISTED_PAGE_SIZE}",
                    headers=headers,
                    timeout=LISTED_REQUEST_TIMEOUT,
                )
                r.raise_for_status()
                payload = r.json()
                break
            except Exception:
                if attempt == LISTED_MAX_RETRIES:
                    raise
                time.sleep(min(2 ** (attempt - 1), 8))

        if payload is None:
            raise ValueError("Brak odpowiedzi payload podczas pobierania listed")

        for row in payload["data"]:
            oid = str(row.get("offer_id", "")).strip()
            try:
                mul = float(row.get("multiplier", 1.0))
            except (TypeError, ValueError):
                mul = 1.0
            if oid:
                result[oid] = mul

        if page >= payload["total_pages"]:
            break
        page += 1

    return result


def _load_multipliers_from_listed_csv(csv_path: Path) -> dict[str, float]:
    listed_df = pd.read_csv(csv_path, sep=";", dtype={"offer_id": str})
    result: dict[str, float] = {}
    for _, row in listed_df.iterrows():
        oid = str(row.get("offer_id", "")).strip()
        if not oid:
            continue
        try:
            mul = float(row.get("multiplier", 1.0))
        except (TypeError, ValueError):
            mul = 1.0
        result[oid] = mul
    return result


class OffersCreatorApp(ctk.CTk):
    def __init__(self):
        super().__init__()

        ctk.set_appearance_mode("light")
        ctk.set_widget_scaling(0.8)
        ctk.set_window_scaling(0.8)

        self.title(APP_TITLE)
        self.geometry("980x760")
        self.minsize(920, 680)
        self.configure(fg_color=APP_BG)

        self.xlsx_path: Path | None = None
        self.listed_csv_path: Path | None = DEFAULT_LISTED_FILE_PATH if DEFAULT_LISTED_FILE_PATH.exists() else None
        self.payloads: list[dict[str, Any]] = []
        self.prepared_requests: list[dict[str, Any]] = []
        self.columns: list[str] = []
        self._stop_event = threading.Event()
        self.api_settings, self.api_settings_path = BolApiConfigStore.load()

        self._build_ui()
        self._apply_api_settings_to_ui()
        self._refresh_spec_labels()
        if self.listed_csv_path:
            self.listed_file_entry.insert(0, str(self.listed_csv_path))

    def _build_ui(self) -> None:
        root = ctk.CTkFrame(self, fg_color=PANEL_BG, corner_radius=14)
        root.pack(fill="both", expand=True, padx=16, pady=16)

        title = ctk.CTkLabel(
            root,
            text="Tworzenie ofert BOL z XLSX (zarys MVP)",
            text_color=TEXT,
            font=ctk.CTkFont(size=24, weight="bold"),
        )
        title.pack(anchor="w", padx=18, pady=(16, 4))

        subtitle = ctk.CTkLabel(
            root,
            text="Ładowanie XLSX, budowa payloadów v11 i opcjonalne wysyłanie CREATE/UPDATE do Offers API.",
            text_color=MUTED,
            font=ctk.CTkFont(size=13),
        )
        subtitle.pack(anchor="w", padx=18, pady=(0, 10))

        mode_frame = ctk.CTkFrame(root, fg_color="transparent")
        mode_frame.pack(fill="x", padx=18, pady=(0, 8))
        ctk.CTkLabel(mode_frame, text="Tryb operacji", text_color=TEXT).pack(side="left", padx=(0, 10))
        self.mode_var = ctk.StringVar(value="CREATE")
        self.mode_selector = ctk.CTkSegmentedButton(
            mode_frame,
            values=["CREATE", "UPDATE"],
            variable=self.mode_var,
            command=self._on_mode_change,
            fg_color=INPUT_BG,
            selected_color=ACCENT,
            selected_hover_color=ACCENT_HOVER,
            unselected_color="white",
            unselected_hover_color="#ffe5f1",
            text_color=TEXT,
        )
        self.mode_selector.pack(side="left")

        files_frame = ctk.CTkFrame(root, fg_color="transparent")
        files_frame.pack(fill="x", padx=18, pady=8)

        self.xlsx_entry = self._build_file_row(
            files_frame,
            label="Plik XLSX",
            button_text="Wybierz XLSX",
            row=0,
            command=self.select_xlsx,
        )
        files_frame.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(files_frame, text="Źródło listed (UPDATE)", text_color=TEXT).grid(
            row=1,
            column=0,
            sticky="w",
            padx=(0, 10),
            pady=6,
        )
        self.listed_source_var = ctk.StringVar(value="API")
        self.listed_source_selector = ctk.CTkSegmentedButton(
            files_frame,
            values=["API", "PLIK"],
            variable=self.listed_source_var,
            fg_color=INPUT_BG,
            selected_color=ACCENT,
            selected_hover_color=ACCENT_HOVER,
            unselected_color="white",
            unselected_hover_color="#ffe5f1",
            text_color=TEXT,
        )
        self.listed_source_selector.grid(row=1, column=1, columnspan=2, sticky="w", pady=6)

        ctk.CTkLabel(files_frame, text="Listed endpoint (UPDATE)", text_color=TEXT).grid(row=2, column=0, sticky="w", padx=(0, 10), pady=6)
        self.listed_entry = ctk.CTkEntry(files_frame, fg_color="white", border_color=BORDER, text_color=TEXT)
        self.listed_entry.grid(row=2, column=1, columnspan=2, sticky="ew", pady=6)
        self.listed_entry.insert(0, DEFAULT_LISTED_ENDPOINT)

        self.listed_file_entry = self._build_file_row(
            files_frame,
            label="Plik listed.csv (UPDATE)",
            button_text="Wybierz CSV",
            row=3,
            command=self.select_listed_csv,
        )

        api_frame = ctk.CTkFrame(root, fg_color=INPUT_BG, corner_radius=10)
        api_frame.pack(fill="x", padx=18, pady=(6, 10))
        api_frame.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(api_frame, text="API URL", text_color=TEXT).grid(row=0, column=0, sticky="w", padx=10, pady=10)
        self.api_entry = ctk.CTkEntry(api_frame, fg_color="white", border_color=BORDER, text_color=TEXT)
        self.api_entry.grid(row=0, column=1, sticky="ew", padx=8, pady=10)
        self.api_entry.insert(0, DEFAULT_API_URL)

        ctk.CTkLabel(api_frame, text="Auth URL", text_color=TEXT).grid(row=1, column=0, sticky="w", padx=10, pady=(0, 10))
        self.auth_entry = ctk.CTkEntry(api_frame, fg_color="white", border_color=BORDER, text_color=TEXT)
        self.auth_entry.grid(row=1, column=1, sticky="ew", padx=8, pady=(0, 10))
        self.auth_entry.insert(0, DEFAULT_AUTH_URL)

        ctk.CTkLabel(api_frame, text="Client ID", text_color=TEXT).grid(row=2, column=0, sticky="w", padx=10, pady=(0, 10))
        self.client_id_entry = ctk.CTkEntry(api_frame, fg_color="white", border_color=BORDER, text_color=TEXT)
        self.client_id_entry.grid(row=2, column=1, sticky="ew", padx=8, pady=(0, 10))

        ctk.CTkLabel(api_frame, text="Client Secret", text_color=TEXT).grid(row=3, column=0, sticky="w", padx=10, pady=(0, 10))
        self.client_secret_entry = ctk.CTkEntry(api_frame, fg_color="white", border_color=BORDER, text_color=TEXT, show="*")
        self.client_secret_entry.grid(row=3, column=1, sticky="ew", padx=8, pady=(0, 10))

        ctk.CTkLabel(api_frame, text="Economic Operators API URL", text_color=TEXT).grid(
            row=4,
            column=0,
            sticky="w",
            padx=10,
            pady=(0, 10),
        )
        self.economic_operators_api_entry = ctk.CTkEntry(api_frame, fg_color="white", border_color=BORDER, text_color=TEXT)
        self.economic_operators_api_entry.grid(row=4, column=1, sticky="ew", padx=8, pady=(0, 10))
        self.economic_operators_api_entry.insert(0, DEFAULT_ECONOMIC_OPERATORS_API_URL)

        self.auto_lookup_offer_id_var = ctk.BooleanVar(value=True)
        self.auto_lookup_offer_id_chk = ctk.CTkCheckBox(
            api_frame,
            text="UPDATE: automatycznie pobieraj offer-id po reference/EAN",
            variable=self.auto_lookup_offer_id_var,
            text_color=TEXT,
            fg_color=ACCENT,
            hover_color=ACCENT_HOVER,
        )
        self.auto_lookup_offer_id_chk.grid(row=5, column=0, columnspan=2, sticky="w", padx=10, pady=(0, 10))

        self.save_api_settings_btn = ctk.CTkButton(
            api_frame,
            text="Zapisz dane API",
            command=self.save_api_settings,
            fg_color="#ffd6e8",
            hover_color="#ffc0dc",
            text_color=TEXT,
            height=30,
            width=180,
        )
        self.save_api_settings_btn.grid(row=6, column=0, columnspan=2, sticky="w", padx=10, pady=(0, 10))

        self.required_label = ctk.CTkLabel(root, text="Wymagane pola z YAML: —", text_color=MUTED, wraplength=900, justify="left")
        self.required_label.pack(anchor="w", padx=18, pady=(0, 8))

        self.fixed_rules_label = ctk.CTkLabel(
            root,
            text="Stałe reguły CREATE: condition = NEW | fulfilment.method = FBR | schedule = BOL_DELIVERY_PROMISE | kraje: NL + BE",
            text_color=MUTED,
            wraplength=900,
            justify="left",
        )
        self.fixed_rules_label.pack(anchor="w", padx=18, pady=(0, 8))

        map_frame = ctk.CTkFrame(root, fg_color=INPUT_BG, corner_radius=10)
        map_frame.pack(fill="x", padx=18, pady=8)
        map_frame.grid_columnconfigure((1, 3), weight=1)

        self.map_entries: dict[str, ctk.CTkEntry] = {}
        mapping_fields = [
            ("offer_id", "Kolumna offer-id (dla UPDATE)"),
            ("ean", "Kolumna EAN"),
            ("price", "Kolumna cena"),
            ("stock", "Kolumna stan"),
            ("reference", "Kolumna reference"),
            ("on_hold", "Kolumna on_hold"),
            ("economic_operator_id", "Kolumna economicOperatorId"),
        ]

        defaults = ColumnMapping()
        for index, (key, caption) in enumerate(mapping_fields):
            row = index // 2
            base_col = (index % 2) * 2
            ctk.CTkLabel(map_frame, text=caption, text_color=TEXT).grid(row=row, column=base_col, sticky="w", padx=10, pady=8)
            entry = ctk.CTkEntry(map_frame, fg_color="white", border_color=BORDER, text_color=TEXT)
            entry.grid(row=row, column=base_col + 1, sticky="ew", padx=8, pady=8)
            entry.insert(0, getattr(defaults, key))
            self.map_entries[key] = entry

        action_row = ctk.CTkFrame(root, fg_color="transparent")
        action_row.pack(fill="x", padx=18, pady=(10, 8))

        self.btn_load_xlsx = self._action_button(action_row, "Wczytaj XLSX", self.load_xlsx_preview)
        self.btn_save = self._action_button(action_row, "Zapisz JSON", self.save_payloads)
        self.btn_send = self._action_button(action_row, "One click processing (XLSX -> API)", self.send_payloads)
        self.btn_stop = ctk.CTkButton(
            action_row,
            text="Zatrzymaj",
            command=self.stop_processing,
            fg_color="#c0392b",
            hover_color="#922b21",
            text_color="white",
            height=36,
        )
        self.btn_stop.pack(side="left", padx=(0, 10))

        self.progress_bar = ctk.CTkProgressBar(root, progress_color=ACCENT)
        self.progress_bar.pack(fill="x", padx=18, pady=(6, 6))
        self.progress_bar.set(0)

        self.status_label = ctk.CTkLabel(root, text="Gotowe.", text_color=MUTED)
        self.status_label.pack(anchor="w", padx=18, pady=(0, 8))

        self.log_box = ctk.CTkTextbox(root, fg_color="white", border_color=BORDER, text_color=TEXT)
        self.log_box.pack(fill="both", expand=True, padx=18, pady=(0, 16))

        if DEFAULT_TEMPLATE_XLSX_PATH.exists():
            self.xlsx_path = DEFAULT_TEMPLATE_XLSX_PATH
            self.xlsx_entry.insert(0, str(DEFAULT_TEMPLATE_XLSX_PATH))

    def _on_mode_change(self, selected_mode: str) -> None:
        self._refresh_api_url_for_mode(selected_mode)
        self._refresh_spec_labels()
        if selected_mode == "UPDATE":
            self.fixed_rules_label.configure(
                text=(
                    "Reguły UPDATE: PATCH wysyła tylko pola obecne w XLSX (brakujące wartości nie są nadpisywane)"
                )
            )
        else:
            self.fixed_rules_label.configure(
                text=(
                    "Stałe reguły CREATE: condition = NEW | fulfilment.method = FBR | "
                    "schedule = BOL_DELIVERY_PROMISE | kraje: NL + BE"
                )
            )

    def _refresh_api_url_for_mode(self, mode: str) -> None:
        current_url = self.api_entry.get().strip()
        default_create = self.api_settings.offers_url or DEFAULT_API_URL
        default_update = self.api_settings.update_offer_url or DEFAULT_UPDATE_API_URL

        if mode == "UPDATE" and current_url in {DEFAULT_API_URL, default_create}:
            self.api_entry.delete(0, "end")
            self.api_entry.insert(0, default_update)
        elif mode == "CREATE" and current_url in {DEFAULT_UPDATE_API_URL, default_update}:
            self.api_entry.delete(0, "end")
            self.api_entry.insert(0, default_create)

    def _refresh_spec_labels(self) -> None:
        mode = self.mode_var.get()
        if mode == "UPDATE":
            fields = ", ".join(UPDATE_PATCH_FIELDS)
            self.required_label.configure(text=f"UPDATE (PatchOfferRequest) — pola dostępne: {fields}")
        else:
            required = ", ".join(CREATE_REQUIRED_FIELDS)
            self.required_label.configure(text=f"CREATE (CreateOfferRequest) — wymagane: {required}")

    def _build_file_row(self, parent: ctk.CTkFrame, label: str, button_text: str, row: int, command) -> ctk.CTkEntry:
        parent.grid_columnconfigure(1, weight=1)
        ctk.CTkLabel(parent, text=label, text_color=TEXT).grid(row=row, column=0, sticky="w", padx=(0, 10), pady=6)
        entry = ctk.CTkEntry(parent, fg_color="white", border_color=BORDER, text_color=TEXT)
        entry.grid(row=row, column=1, sticky="ew", pady=6)
        ctk.CTkButton(
            parent,
            text=button_text,
            command=command,
            fg_color=ACCENT,
            hover_color=ACCENT_HOVER,
            text_color="white",
            width=140,
        ).grid(row=row, column=2, sticky="e", padx=(10, 0), pady=6)
        return entry

    def _action_button(self, parent: ctk.CTkFrame, text: str, command) -> ctk.CTkButton:
        button = ctk.CTkButton(
            parent,
            text=text,
            command=command,
            fg_color=ACCENT,
            hover_color=ACCENT_HOVER,
            text_color="white",
            height=36,
        )
        button.pack(side="left", padx=(0, 10))
        return button

    def _apply_api_settings_to_ui(self) -> None:
        self.client_id_entry.delete(0, "end")
        self.client_id_entry.insert(0, self.api_settings.client_id)
        self.client_secret_entry.delete(0, "end")
        self.client_secret_entry.insert(0, self.api_settings.client_secret)
        self.auth_entry.delete(0, "end")
        self.auth_entry.insert(0, self.api_settings.auth_url or DEFAULT_AUTH_URL)
        self.economic_operators_api_entry.delete(0, "end")
        self.economic_operators_api_entry.insert(
            0,
            self.api_settings.economic_operators_url or DEFAULT_ECONOMIC_OPERATORS_API_URL,
        )

        default_api = self.api_settings.offers_url or DEFAULT_API_URL
        self.api_entry.delete(0, "end")
        self.api_entry.insert(0, default_api)

    def _collect_api_settings_from_ui(self) -> BolApiSettings:
        current_mode = self.mode_var.get()
        api_url = self.api_entry.get().strip()

        offers_url = self.api_settings.offers_url or DEFAULT_API_URL
        update_offer_url = self.api_settings.update_offer_url or DEFAULT_UPDATE_API_URL
        if current_mode == "UPDATE":
            update_offer_url = api_url or DEFAULT_UPDATE_API_URL
        else:
            offers_url = api_url or DEFAULT_API_URL

        return BolApiSettings(
            client_id=self.client_id_entry.get().strip(),
            client_secret=self.client_secret_entry.get().strip(),
            auth_url=self.auth_entry.get().strip() or DEFAULT_AUTH_URL,
            offers_url=offers_url,
            update_offer_url=update_offer_url,
            economic_operators_url=self.economic_operators_api_entry.get().strip() or DEFAULT_ECONOMIC_OPERATORS_API_URL,
        )

    def save_api_settings(self) -> None:
        try:
            self.api_settings = self._collect_api_settings_from_ui()
            self.api_settings_path = BolApiConfigStore.save(self.api_settings, self.api_settings_path)
            self._log(f"Zapisano dane API do: {self.api_settings_path}")
            self._set_status("Dane API zapisane.")
        except Exception as exc:
            self._log(f"Błąd zapisu danych API: {exc}")
            self._set_status(f"Błąd zapisu danych API: {exc}", is_error=True)

    def _log(self, message: str) -> None:
        timestamp = datetime.now().strftime("%H:%M:%S")

        def _update() -> None:
            self.log_box.insert("end", f"[{timestamp}] {message}\n")
            self.log_box.see("end")

        self.after(0, _update)

    def _set_status(self, text: str, is_error: bool = False) -> None:
        self.after(0, lambda: self.status_label.configure(text=text, text_color=ERROR if is_error else SUCCESS))

    def _set_progress(self, value: float) -> None:
        safe_value = min(max(value, 0.0), 1.0)
        self.after(0, lambda: self.progress_bar.set(safe_value))

    def stop_processing(self) -> None:
        self._stop_event.set()
        self._log("Zatrzymywanie — czekaj na przerwanie bieżącej operacji...")
        self._set_status("Zatrzymywanie...", is_error=True)

    def select_xlsx(self) -> None:
        path = filedialog.askopenfilename(
            title="Wybierz plik XLSX",
            filetypes=[("Excel", "*.xlsx *.xls"), ("All files", "*.*")],
        )
        if not path:
            return
        self.xlsx_path = Path(path)
        self.xlsx_entry.delete(0, "end")
        self.xlsx_entry.insert(0, str(self.xlsx_path))
        self._log(f"Wybrano XLSX: {self.xlsx_path}")

    def select_listed_csv(self) -> None:
        path = filedialog.askopenfilename(
            title="Wybierz listed.csv",
            filetypes=[("CSV", "*.csv"), ("All files", "*.*")],
        )
        if not path:
            return
        self.listed_csv_path = Path(path)
        self.listed_file_entry.delete(0, "end")
        self.listed_file_entry.insert(0, str(self.listed_csv_path))
        self._log(f"Wybrano listed.csv: {self.listed_csv_path}")

    def load_xlsx_preview(self) -> None:
        if not self.xlsx_path or not self.xlsx_path.exists():
            messagebox.showwarning("Brak XLSX", "Wybierz plik XLSX.")
            return

        try:
            df = pd.read_excel(self.xlsx_path)
            self.columns = [str(column) for column in df.columns]
            self._log(f"Załadowano XLSX: {self.xlsx_path.name} | wiersze: {len(df)} | kolumny: {len(self.columns)}")
            self._log(f"Kolumny: {', '.join(self.columns[:20])}{' ...' if len(self.columns) > 20 else ''}")
            self._set_status("XLSX załadowany poprawnie.")
        except Exception as exc:
            self._set_status(f"Błąd XLSX: {exc}", is_error=True)
            self._log(f"Błąd odczytu XLSX: {exc}")

    def _build_mapping(self) -> ColumnMapping:
        return ColumnMapping(
            offer_id=self.map_entries["offer_id"].get().strip(),
            ean=self.map_entries["ean"].get().strip(),
            price=self.map_entries["price"].get().strip(),
            stock=self.map_entries["stock"].get().strip(),
            reference=self.map_entries["reference"].get().strip(),
            on_hold=self.map_entries["on_hold"].get().strip(),
            economic_operator_id=self.map_entries["economic_operator_id"].get().strip(),
        )

    def _generate_payloads_internal(
        self,
        preview_prices: bool = True,
        preview_first_payload: bool = True,
        require_offer_id_for_update: bool = True,
    ) -> bool:
        if not self.xlsx_path or not self.xlsx_path.exists():
            messagebox.showwarning("Brak XLSX", "Wybierz plik XLSX.")
            return False

        try:
            df = pd.read_excel(self.xlsx_path)
        except Exception as exc:
            self._set_status(f"Nie można odczytać XLSX: {exc}", is_error=True)
            return False

        mapping = self._build_mapping()
        defaults = OfferDefaults()
        mode = self.mode_var.get()

        offer_id_to_multiplier: dict[str, float] = {}
        if mode == "UPDATE":
            listed_source = self.listed_source_var.get().strip().upper() or "API"
            if listed_source == "PLIK":
                listed_path = Path(self.listed_file_entry.get().strip()) if self.listed_file_entry.get().strip() else None
                if listed_path and listed_path.exists():
                    try:
                        offer_id_to_multiplier = _load_multipliers_from_listed_csv(listed_path)
                        self._log(
                            f"Wczytano listed z pliku: {len(offer_id_to_multiplier)} mnożników ({listed_path.name})"
                        )
                    except Exception as exc:
                        self._log(f"Ostrzeżenie: nie można wczytać listed z pliku ({exc}), mnożniki = 1.0")
                else:
                    self._log("Ostrzeżenie: brak poprawnego pliku listed.csv, mnożniki = 1.0")
            else:
                endpoint = self.listed_entry.get().strip() or DEFAULT_LISTED_ENDPOINT
                try:
                    offer_id_to_multiplier = _fetch_listed_multipliers(endpoint)
                    self._log(
                        f"Pobrano listed z API: {len(offer_id_to_multiplier)} mnożników (endpoint: {endpoint})"
                    )
                except Exception as exc:
                    self._log(f"Ostrzeżenie: nie można pobrać listed z API ({exc}), mnożniki = 1.0")

        payloads: list[dict[str, Any]] = []
        prepared_requests: list[dict[str, Any]] = []
        errors = 0
        price_preview_count = 0

        total = len(df)
        self._set_progress(0)

        for index, row in enumerate(df.to_dict(orient="records"), start=1):
            try:
                if mode == "UPDATE":
                    row_offer_id = OfferPayloadBuilder._normalize_offer_id(row.get(mapping.offer_id, ""))
                    multiplier = offer_id_to_multiplier.get(row_offer_id, 1.0)
                    offer_id, payload = OfferPayloadBuilder.build_update_request(
                        row,
                        mapping,
                        defaults,
                        require_offer_id=require_offer_id_for_update,
                        multiplier=multiplier,
                    )
                else:
                    offer_id = None
                    payload = OfferPayloadBuilder.build_create_payload(row, mapping, defaults)

                payloads.append(payload)
                prepared_requests.append(
                    {
                        "offer_id": offer_id,
                        "payload": payload,
                        "reference": str(row.get(mapping.reference, "") or "").strip(),
                        "ean": OfferPayloadBuilder._normalize_ean(row.get(mapping.ean, "")),
                        "row_index": index - 1,
                    }
                )

                if preview_prices and price_preview_count < 5:
                    raw_price_value = row.get(mapping.price)
                    raw_price = (
                        OfferPayloadBuilder._to_float(raw_price_value, fallback=0.0)
                        if OfferPayloadBuilder._has_value(raw_price_value)
                        else 0.0
                    )
                    bundle_prices = payload.get("pricing", {}).get("bundlePrices", [])
                    if raw_price > 0 and bundle_prices:
                        tiers_text = " | ".join(
                            f"q{tier.get('quantity')}={tier.get('unitPrice')}" for tier in bundle_prices
                        )
                        self._log(f"Podgląd ceny: wejście={raw_price} -> {tiers_text}")
                        price_preview_count += 1
            except Exception as exc:
                errors += 1
                if errors <= 10:
                    self._log(f"Wiersz {index}: pominięty ({exc})")

            if total > 0:
                self._set_progress(index / total)

            if self._stop_event.is_set():
                self._log("Przerwano przez użytkownika podczas generowania payloadów.")
                break

        self.payloads = payloads
        self.prepared_requests = prepared_requests
        self._log(f"Tryb {mode}: wygenerowano rekordów: {len(payloads)} | pominięte: {errors}")

        if preview_first_payload and payloads:
            first_request = prepared_requests[0]
            if mode == "UPDATE":
                preview_payload = {
                    "offerId": first_request.get("offer_id"),
                    "payload": first_request.get("payload"),
                }
            else:
                preview_payload = payloads[0]
            preview = json.dumps(preview_payload, ensure_ascii=False, indent=2)
            self._log("Podgląd pierwszego payloadu:")
            self._log(preview)

        if not payloads:
            self._set_status("Brak poprawnych payloadów do przetworzenia.", is_error=True)
            return False

        if errors > 0:
            self._set_status(f"Gotowe z ostrzeżeniami. Payloady: {len(payloads)} | Błędy: {errors}", is_error=True)
        else:
            self._set_status(f"Gotowe. Wygenerowano {len(payloads)} payloadów.")

        return True

    def generate_payloads(self) -> None:
        self._generate_payloads_internal(preview_prices=True, preview_first_payload=True)

    def save_payloads(self) -> None:
        if not self.payloads:
            self._log("Brak payloadów w pamięci. Automatyczne generowanie z XLSX przed zapisem JSON.")
            if not self._generate_payloads_internal(preview_prices=False, preview_first_payload=False):
                return

        mode = self.mode_var.get()
        data_to_save: Any = self.prepared_requests if mode == "UPDATE" else self.payloads
        default_name = "offers_update_payloads.json" if mode == "UPDATE" else "offers_payloads.json"

        output_path = filedialog.asksaveasfilename(
            title="Zapisz payloady JSON",
            defaultextension=".json",
            filetypes=[("JSON", "*.json")],
            initialfile=default_name,
        )
        if not output_path:
            return

        try:
            with open(output_path, "w", encoding="utf-8") as handle:
                json.dump(data_to_save, handle, ensure_ascii=False, indent=2)
            self._log(f"Zapisano plik JSON: {output_path}")
            self._set_status("Payloady zapisane.")
        except Exception as exc:
            self._set_status(f"Błąd zapisu JSON: {exc}", is_error=True)

    def send_payloads(self) -> None:
        mode = self.mode_var.get()
        allow_auto_lookup = mode == "UPDATE" and bool(self.auto_lookup_offer_id_var.get())
        self._log(f"One click processing ({mode}): start automatycznego budowania payloadów z XLSX.")
        self._stop_event.clear()

        client_id = self.client_id_entry.get().strip()
        client_secret = self.client_secret_entry.get().strip()
        if not client_id or not client_secret:
            messagebox.showwarning("Brak danych API", "Wpisz Client ID i Client Secret.")
            return

        api_url = self.api_entry.get().strip() or DEFAULT_API_URL
        auth_url = self.auth_entry.get().strip() or DEFAULT_AUTH_URL
        economic_operators_api_url = self.economic_operators_api_entry.get().strip() or DEFAULT_ECONOMIC_OPERATORS_API_URL
        self.save_api_settings()
        mapping = self._build_mapping()

        def _lookup_then_send() -> None:
            if not self._generate_payloads_internal(
                preview_prices=False,
                preview_first_payload=False,
                require_offer_id_for_update=not allow_auto_lookup,
            ):
                return

            if self._stop_event.is_set():
                self._set_status("Zatrzymano.", is_error=True)
                return

            if allow_auto_lookup:
                missing_before = sum(1 for item in self.prepared_requests if not str(item.get("offer_id") or "").strip())
                if missing_before > 0:
                    self._log(
                        f"UPDATE: brakujących offer-id: {missing_before}. Start automatycznego lookup po reference/EAN."
                    )
                    resolved, unresolved = self._resolve_missing_offer_ids(api_url, auth_url, client_id, client_secret)
                    self._log(f"Lookup zakończony. Uzupełniono: {resolved} | nadal brak: {unresolved}")
                    if unresolved > 0:
                        self._set_status(
                            f"Nie znaleziono {unresolved} offer-id (sprawdź reference/EAN).", is_error=True
                        )

            if self._stop_event.is_set():
                self._set_status("Zatrzymano.", is_error=True)
                return

            self._log(f"{mode}: mapowanie economicOperatorId (nazwa -> UUID) przez Economic Operators API.")
            eo_resolved, eo_unresolved = self._resolve_economic_operator_ids(
                economic_operators_api_url,
                auth_url,
                client_id,
                client_secret,
            )
            if eo_resolved or eo_unresolved:
                self._log(f"EconomicOperator mapowanie: uzupełnione={eo_resolved} | nieudane={eo_unresolved}")

            if self._stop_event.is_set():
                self._set_status("Zatrzymano.", is_error=True)
                return

            if mode == "UPDATE":
                self._log("UPDATE: weryfikacja offer-id względem reference/EAN przed wysyłką.")
                verified, rejected = self._verify_update_offer_targets(api_url, auth_url, client_id, client_secret)
                self._log(f"Weryfikacja UPDATE: poprawne={verified} | odrzucone={rejected}")
                if verified == 0:
                    self._set_status("Brak poprawnych rekordów po weryfikacji reference/EAN.", is_error=True)
                    return

            if self._stop_event.is_set():
                self._set_status("Zatrzymano.", is_error=True)
                return

            self._send_worker(mode, api_url, auth_url, client_id, client_secret, mapping)

        worker = threading.Thread(target=_lookup_then_send, daemon=True)
        worker.start()

    @staticmethod
    def _build_update_url(api_url: str, offer_id: str) -> str:
        url = api_url.strip()
        if "{offer-id}" in url:
            return url.replace("{offer-id}", offer_id)
        if "{offer_id}" in url:
            return url.replace("{offer_id}", offer_id)
        return f"{url.rstrip('/')}/{offer_id}"

    @staticmethod
    def _build_offers_collection_url(api_url: str) -> str:
        url = api_url.strip()
        if "{offer-id}" in url:
            return url.replace("/{offer-id}", "").replace("{offer-id}", "")
        if "{offer_id}" in url:
            return url.replace("/{offer_id}", "").replace("{offer_id}", "")
        if url.rstrip("/").endswith("/offers"):
            return url.rstrip("/")
        if "/offers/" in url:
            return url.rsplit("/", 1)[0]
        return DEFAULT_API_URL

    @staticmethod
    def _looks_like_uuid(value: str) -> bool:
        return bool(re.fullmatch(r"[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}", value))

    def _find_economic_operator_id_by_name(
        self,
        economic_operators_url: str,
        token_manager: OAuthTokenManager,
        limiter: RateLimiter,
        operator_name: str,
    ) -> str | None:
        query = urlencode({"name": operator_name, "page-size": 100})
        request_url = f"{economic_operators_url.rstrip('/')}?{query}"

        for attempt in range(1, SEND_MAX_RETRIES + 1):
            limiter.acquire()
            try:
                headers = {
                    "Accept": "application/vnd.economic-operator.v1+json",
                    "Authorization": f"Bearer {token_manager.get_token()}",
                }
                response = requests.get(request_url, headers=headers, timeout=LOOKUP_REQUEST_TIMEOUT)

                if response.status_code == 200:
                    operators = response.json().get("operators", [])
                    if not operators:
                        return None

                    lowered_name = operator_name.strip().lower()
                    exact = [
                        item for item in operators if str(item.get("name", "")).strip().lower() == lowered_name
                    ]
                    candidates = exact if exact else operators
                    if len(candidates) > 1:
                        self._log(
                            f"EconomicOperator lookup: wiele wyników dla '{operator_name}', używam pierwszego z {len(candidates)}"
                        )

                    economic_operator_id = str(candidates[0].get("id", "")).strip()
                    return economic_operator_id or None

                if response.status_code == 429 and attempt < SEND_MAX_RETRIES:
                    retry_after_raw = response.headers.get("Retry-After")
                    try:
                        retry_after = int(float(retry_after_raw)) if retry_after_raw is not None else 1
                    except (TypeError, ValueError):
                        retry_after = 1
                    time.sleep(max(retry_after, 1))
                    continue

                if 500 <= response.status_code <= 599 and attempt < SEND_MAX_RETRIES:
                    time.sleep(min(2 ** attempt, 8))
                    continue

                self._log(
                    f"EconomicOperator lookup: błąd HTTP {response.status_code} dla '{operator_name}': {response.text[:250]}"
                )
                return None
            except Exception as exc:
                if attempt < SEND_MAX_RETRIES:
                    time.sleep(min(2 ** attempt, 8))
                    continue
                self._log(f"EconomicOperator lookup: wyjątek dla '{operator_name}': {exc}")
                return None

        return None

    def _resolve_economic_operator_ids(
        self,
        economic_operators_api_url: str,
        auth_url: str,
        client_id: str,
        client_secret: str,
    ) -> tuple[int, int]:
        token_manager = OAuthTokenManager(auth_url=auth_url, client_id=client_id, client_secret=client_secret)
        limiter = RateLimiter(LOOKUP_RATE_LIMIT_PER_SEC)

        name_to_indices: dict[str, list[int]] = {}
        for i, req in enumerate(self.prepared_requests):
            payload = req.get("payload", {})
            raw_value = str(payload.get("economicOperatorId", "") or "").strip()
            if not raw_value:
                continue
            if self._looks_like_uuid(raw_value):
                continue
            name_to_indices.setdefault(raw_value, []).append(i)

        if not name_to_indices:
            return 0, 0

        self._log(
            f"EconomicOperator lookup: {len(name_to_indices)} unikalnych nazw do zamiany na UUID (workers={LOOKUP_MAX_WORKERS})"
        )

        resolved = 0
        unresolved = 0
        with ThreadPoolExecutor(max_workers=LOOKUP_MAX_WORKERS) as executor:
            futures = {
                executor.submit(
                    self._find_economic_operator_id_by_name,
                    economic_operators_api_url,
                    token_manager,
                    limiter,
                    operator_name,
                ): operator_name
                for operator_name in name_to_indices.keys()
            }

            for future in as_completed(futures):
                operator_name = futures[future]
                found_id = future.result()
                indices = name_to_indices[operator_name]
                if found_id and self._looks_like_uuid(found_id):
                    for idx in indices:
                        self.prepared_requests[idx]["payload"]["economicOperatorId"] = found_id
                    resolved += len(indices)
                    self._log(f"EconomicOperator OK: '{operator_name}' -> {found_id} ({len(indices)} wierszy)")
                else:
                    unresolved += len(indices)
                    self._log(f"EconomicOperator FAIL: '{operator_name}'")

        return resolved, unresolved

    def _find_offer_id_for_row(
        self,
        offers_url: str,
        token_manager: OAuthTokenManager,
        limiter: RateLimiter,
        reference: str,
        ean: str,
    ) -> str | None:
        def _query(params: dict[str, Any], match_key: str, match_value: str) -> str | None:
            for attempt in range(1, SEND_MAX_RETRIES + 1):
                limiter.acquire()
                try:
                    headers = {
                        "Accept": "application/vnd.retailer.v11+json",
                        "Authorization": f"Bearer {token_manager.get_token()}",
                    }
                    response = requests.get(
                        offers_url,
                        headers=headers,
                        params=params,
                        timeout=LOOKUP_REQUEST_TIMEOUT,
                    )
                    if response.status_code == 200:
                        offers = response.json().get("offers", [])
                        exact_matches = [item for item in offers if str(item.get(match_key, "")).strip() == match_value]
                        candidates = exact_matches if exact_matches else offers
                        if not candidates:
                            return None

                        if reference:
                            ref_matches = [item for item in candidates if str(item.get("reference", "")).strip() == reference]
                            if ref_matches:
                                candidates = ref_matches

                        if ean:
                            ean_matches = [item for item in candidates if str(item.get("ean", "")).strip() == ean]
                            if ean_matches:
                                candidates = ean_matches

                        if len(candidates) > 1:
                            self._log(
                                f"Lookup: wiele ofert dla {match_key}={match_value}; używam pierwszej z {len(candidates)}"
                            )
                        offer_id_value = str(candidates[0].get("offerId", "")).strip()
                        return offer_id_value or None

                    if response.status_code == 429 and attempt < SEND_MAX_RETRIES:
                        retry_after_raw = response.headers.get("Retry-After")
                        try:
                            retry_after = int(float(retry_after_raw)) if retry_after_raw is not None else 1
                        except (TypeError, ValueError):
                            retry_after = 1
                        time.sleep(max(retry_after, 1))
                        continue

                    if 500 <= response.status_code <= 599 and attempt < SEND_MAX_RETRIES:
                        time.sleep(min(2 ** attempt, 8))
                        continue

                    self._log(
                        f"Lookup: błąd HTTP {response.status_code} dla {match_key}={match_value}: {response.text[:250]}"
                    )
                    return None
                except Exception as exc:
                    if attempt < SEND_MAX_RETRIES:
                        time.sleep(min(2 ** attempt, 8))
                        continue
                    self._log(f"Lookup: wyjątek dla {match_key}={match_value}: {exc}")
                    return None

            return None

        if reference:
            offer_id_by_reference = _query(
                params={"reference": reference, "page-size": 100},
                match_key="reference",
                match_value=reference,
            )
            if offer_id_by_reference:
                return offer_id_by_reference

        if ean:
            offer_id_by_ean = _query(
                params={"eans": [ean], "page-size": 100},
                match_key="ean",
                match_value=ean,
            )
            if offer_id_by_ean:
                return offer_id_by_ean

        return None

    def _resolve_missing_offer_ids(
        self,
        api_url: str,
        auth_url: str,
        client_id: str,
        client_secret: str,
        only_successful: bool = False,
    ) -> tuple[int, int]:
        offers_url = self._build_offers_collection_url(api_url)
        token_manager = OAuthTokenManager(auth_url=auth_url, client_id=client_id, client_secret=client_secret)
        limiter = RateLimiter(LOOKUP_RATE_LIMIT_PER_SEC)

        # Group rows by (reference, ean) to avoid duplicate HTTP calls
        key_to_indices: dict[tuple[str, str], list[int]] = {}
        no_identifier_indices: list[int] = []
        considered = 0
        for i, req in enumerate(self.prepared_requests):
            if only_successful and not bool(req.get("send_success")):
                continue
            if str(req.get("offer_id") or "").strip():
                continue

            considered += 1
            payload = req.get("payload", {})
            reference = str(req.get("reference") or payload.get("reference") or "").strip()
            ean = str(req.get("ean") or "").strip()
            if not reference and not ean:
                no_identifier_indices.append(i)
                continue
            key = (reference, ean)
            key_to_indices.setdefault(key, []).append(i)

        unique_keys = list(key_to_indices.keys())
        self._log(
            f"Lookup: {considered} wierszy -> {len(unique_keys)} unikalnych zapytań (workers={LOOKUP_MAX_WORKERS})"
        )

        if not unique_keys:
            unresolved_without_identifier = len(no_identifier_indices)
            if unresolved_without_identifier:
                self._log(
                    f"Lookup: pominięto {unresolved_without_identifier} wierszy bez reference i EAN (nie można pobrać offer-id)."
                )
            return 0, unresolved_without_identifier

        results: dict[tuple[str, str], str | None] = {}

        def _lookup_key(key: tuple[str, str]) -> tuple[tuple[str, str], str | None]:
            reference, ean = key
            found = self._find_offer_id_for_row(
                offers_url=offers_url,
                token_manager=token_manager,
                limiter=limiter,
                reference=reference,
                ean=ean,
            )
            return key, found

        with ThreadPoolExecutor(max_workers=LOOKUP_MAX_WORKERS) as executor:
            futures = {executor.submit(_lookup_key, key): key for key in unique_keys}
            for future in as_completed(futures):
                key, found_offer_id = future.result()
                results[key] = found_offer_id

        resolved = 0
        unresolved = len(no_identifier_indices)
        if no_identifier_indices:
            self._log(
                f"Lookup: pominięto {len(no_identifier_indices)} wierszy bez reference i EAN (nie można pobrać offer-id)."
            )

        for key, indices in key_to_indices.items():
            found_offer_id = results.get(key)
            reference, ean = key
            if found_offer_id:
                for i in indices:
                    self.prepared_requests[i]["offer_id"] = found_offer_id
                resolved += len(indices)
                self._log(f"Lookup OK: offer-id={found_offer_id} | reference={reference or '-'} | ean={ean or '-'}")
            else:
                unresolved += len(indices)
                self._log(f"Lookup FAIL: reference={reference or '-'} | ean={ean or '-'}")

        return resolved, unresolved

    @staticmethod
    def _extract_offer_id_from_create_response(response: requests.Response) -> str:
        try:
            body = response.json()
        except ValueError:
            return ""

        candidate_keys = {"offerId", "offer_id", "offer-id"}

        def _walk(value: Any) -> str:
            if isinstance(value, dict):
                for key in candidate_keys:
                    candidate = value.get(key)
                    if candidate is not None:
                        candidate_str = str(candidate).strip()
                        if candidate_str:
                            return candidate_str
                for nested in value.values():
                    found = _walk(nested)
                    if found:
                        return found
            elif isinstance(value, list):
                for item in value:
                    found = _walk(item)
                    if found:
                        return found
            return ""

        return _walk(body)

    def _persist_offer_ids_to_xlsx(self, mapping: ColumnMapping, only_successful: bool = False) -> tuple[int, int]:
        if not self.xlsx_path or not self.xlsx_path.exists():
            raise ValueError("Brak pliku XLSX do aktualizacji offer-id")

        allowed_suffixes = {".xlsx", ".xlsm", ".xltx", ".xltm"}
        if self.xlsx_path.suffix.lower() not in allowed_suffixes:
            raise ValueError("Automatyczny zapis offer-id obsługuje tylko pliki XLSX/XLSM")

        assignments: list[tuple[int, str]] = []
        missing_offer_id = 0
        for request_data in self.prepared_requests:
            if only_successful and not bool(request_data.get("send_success")):
                continue

            row_index = request_data.get("row_index")
            offer_id = str(request_data.get("offer_id") or "").strip()
            if not offer_id:
                missing_offer_id += 1
                continue

            try:
                excel_row = int(row_index) + 2
            except (TypeError, ValueError):
                missing_offer_id += 1
                continue

            assignments.append((excel_row, offer_id))

        if not assignments:
            return 0, missing_offer_id

        workbook = load_workbook(filename=self.xlsx_path)
        worksheet = workbook.worksheets[0]

        headers: dict[str, int] = {}
        for column_index in range(1, worksheet.max_column + 1):
            header_value = worksheet.cell(row=1, column=column_index).value
            if header_value is None:
                continue
            headers[str(header_value).strip()] = column_index

        offer_id_column = headers.get(mapping.offer_id)
        if not offer_id_column:
            offer_id_column = worksheet.max_column + 1
            worksheet.cell(row=1, column=offer_id_column, value=mapping.offer_id)

        for excel_row, offer_id in assignments:
            worksheet.cell(row=excel_row, column=offer_id_column, value=offer_id)

        workbook.save(self.xlsx_path)
        return len(assignments), missing_offer_id

    def _verify_update_offer_targets(
        self,
        api_url: str,
        auth_url: str,
        client_id: str,
        client_secret: str,
    ) -> tuple[int, int]:
        token_manager = OAuthTokenManager(auth_url=auth_url, client_id=client_id, client_secret=client_secret)
        limiter = RateLimiter(LOOKUP_RATE_LIMIT_PER_SEC)

        def _verify_single(prepared: dict[str, Any]) -> tuple[bool, str]:
            offer_id = str(prepared.get("offer_id") or "").strip()
            if not offer_id:
                return False, "Brak offer-id przed weryfikacją"

            expected_reference = str(prepared.get("reference") or "").strip()
            expected_ean = str(prepared.get("ean") or "").strip()
            if not expected_reference and not expected_ean:
                return False, f"Brak reference/EAN do weryfikacji | offer-id={offer_id}"

            request_url = self._build_update_url(api_url, offer_id)

            for attempt in range(1, SEND_MAX_RETRIES + 1):
                limiter.acquire()
                try:
                    headers = {
                        "Accept": "application/vnd.retailer.v11+json",
                        "Authorization": f"Bearer {token_manager.get_token()}",
                    }
                    response = requests.get(request_url, headers=headers, timeout=LOOKUP_REQUEST_TIMEOUT)

                    if response.status_code == 200:
                        body = response.json()
                        offer = body.get("offer", body)
                        actual_reference = str(offer.get("reference", "") or "").strip()
                        actual_ean = str(offer.get("ean", "") or "").strip()

                        if expected_reference and actual_reference != expected_reference:
                            return (
                                False,
                                f"Reference mismatch | offer-id={offer_id} | oczekiwany={expected_reference} | API={actual_reference or '-'}",
                            )
                        if expected_ean and actual_ean and actual_ean != expected_ean:
                            return (
                                False,
                                f"EAN mismatch | offer-id={offer_id} | oczekiwany={expected_ean} | API={actual_ean}",
                            )

                        return True, f"Weryfikacja OK | offer-id={offer_id}"

                    if response.status_code == 429 and attempt < SEND_MAX_RETRIES:
                        retry_after_raw = response.headers.get("Retry-After")
                        try:
                            retry_after = int(float(retry_after_raw)) if retry_after_raw is not None else 1
                        except (TypeError, ValueError):
                            retry_after = 1
                        time.sleep(max(retry_after, 1))
                        continue

                    if 500 <= response.status_code <= 599 and attempt < SEND_MAX_RETRIES:
                        time.sleep(min(2 ** attempt, 8))
                        continue

                    return False, f"HTTP {response.status_code} podczas weryfikacji offer-id={offer_id}"
                except Exception as exc:
                    if attempt < SEND_MAX_RETRIES:
                        time.sleep(min(2 ** attempt, 8))
                        continue
                    return False, f"Wyjątek weryfikacji offer-id={offer_id}: {exc}"

            return False, f"Weryfikacja nieudana | offer-id={offer_id}"

        if not self.prepared_requests:
            return 0, 0

        verified_requests: list[dict[str, Any]] = []
        rejected = 0

        with ThreadPoolExecutor(max_workers=LOOKUP_MAX_WORKERS) as executor:
            futures = {executor.submit(_verify_single, req): req for req in self.prepared_requests}
            for future in as_completed(futures):
                success, message = future.result()
                self._log(message)
                if success:
                    verified_requests.append(futures[future])
                else:
                    rejected += 1

        self.prepared_requests = verified_requests
        return len(verified_requests), rejected

    def _send_worker(
        self,
        mode: str,
        api_url: str,
        auth_url: str,
        client_id: str,
        client_secret: str,
        mapping: ColumnMapping,
    ) -> None:
        ok_count = 0
        fail_count = 0
        total = len(self.prepared_requests)

        limiter = RateLimiter(SEND_RATE_LIMIT_PER_SEC)
        token_manager = OAuthTokenManager(auth_url=auth_url, client_id=client_id, client_secret=client_secret)
        self._set_progress(0)
        self._log(f"Start wysyłki [{mode}]: {total} ofert -> {api_url} | limit {SEND_RATE_LIMIT_PER_SEC}/s")

        def _send_single(prepared: dict[str, Any]) -> tuple[bool, str, str]:
            payload = prepared.get("payload", {})
            offer_id = str(prepared.get("offer_id") or "").strip()
            offer_ref = str(offer_id or payload.get("reference") or payload.get("ean") or "brak_ref")
            for attempt in range(1, SEND_MAX_RETRIES + 1):
                limiter.acquire()
                try:
                    token = token_manager.get_token()
                    headers = {
                        "Accept": "application/vnd.retailer.v11+json",
                        "Content-Type": "application/vnd.retailer.v11+json",
                        "Authorization": f"Bearer {token}",
                    }
                    if mode == "UPDATE":
                        if not offer_id:
                            return False, "Brak offer-id dla UPDATE"
                        request_url = self._build_update_url(api_url, offer_id)
                        response = requests.patch(request_url, headers=headers, json=payload, timeout=SEND_REQUEST_TIMEOUT)
                        method_name = "PATCH"
                    else:
                        request_url = api_url
                        response = requests.post(request_url, headers=headers, json=payload, timeout=SEND_REQUEST_TIMEOUT)
                        method_name = "POST"

                    response_body = response.text if response.text else "<empty>"
                    self._log(
                        f"API odpowiedź {method_name} | ref={offer_ref} | próba {attempt}/{SEND_MAX_RETRIES} | "
                        f"HTTP {response.status_code} | {response_body}"
                    )

                    expected_codes = (204,) if mode == "UPDATE" else (200, 201, 202)
                    if response.status_code in expected_codes:
                        created_offer_id = ""
                        if mode == "CREATE":
                            created_offer_id = self._extract_offer_id_from_create_response(response)
                            if created_offer_id:
                                return (
                                    True,
                                    f"OK | ref={offer_ref} | HTTP {response.status_code} | offer-id={created_offer_id}",
                                    created_offer_id,
                                )

                        return True, f"OK | ref={offer_ref} | HTTP {response.status_code}", created_offer_id

                    if response.status_code == 429:
                        retry_after_raw = response.headers.get("Retry-After")
                        reset_raw = response.headers.get("x-ratelimit-reset")
                        try:
                            retry_after = int(float(retry_after_raw)) if retry_after_raw is not None else 1
                        except (TypeError, ValueError):
                            retry_after = 1
                        try:
                            reset_after = int(float(reset_raw)) if reset_raw is not None else 0
                        except (TypeError, ValueError):
                            reset_after = 0

                        sleep_seconds = max(retry_after, reset_after + 1, 1)
                        if attempt < SEND_MAX_RETRIES:
                            time.sleep(sleep_seconds)
                            continue
                        return False, f"HTTP 429 po retry | ref={offer_ref}", ""

                    if 500 <= response.status_code <= 599 and attempt < SEND_MAX_RETRIES:
                        time.sleep(min(2 ** attempt, 8))
                        continue

                    return False, f"HTTP {response.status_code} | ref={offer_ref}", ""
                except Exception as exc:
                    exc_text = str(exc)
                    if "token" in exc_text.lower() and attempt < SEND_MAX_RETRIES:
                        time.sleep(1)
                        continue
                    if attempt < SEND_MAX_RETRIES:
                        time.sleep(min(2 ** attempt, 8))
                        continue
                    return False, f"Błąd request | ref={offer_ref}: {exc}", ""

            return False, "Nieznany błąd wysyłki", ""

        completed = 0
        with ThreadPoolExecutor(max_workers=SEND_MAX_WORKERS) as executor:
            futures = {executor.submit(_send_single, request_data): request_data for request_data in self.prepared_requests}

            for future in as_completed(futures):
                completed += 1
                request_data = futures[future]
                success, message, created_offer_id = future.result()
                request_data["send_success"] = success
                if success and mode == "CREATE" and created_offer_id:
                    request_data["offer_id"] = created_offer_id

                self._log(message)
                if success:
                    ok_count += 1
                else:
                    fail_count += 1

                self._set_progress(completed / total if total else 1.0)

                if self._stop_event.is_set():
                    self._log("Przerwano wysyłkę przez użytkownika.")
                    for f in futures:
                        f.cancel()
                    break

        if mode == "CREATE" and ok_count > 0:
            missing_offer_ids = sum(
                1
                for item in self.prepared_requests
                if item.get("send_success") and not str(item.get("offer_id") or "").strip()
            )
            if missing_offer_ids > 0:
                self._log(
                    f"CREATE: brakujących offer-id po odpowiedzi API: {missing_offer_ids}. Lookup po reference/EAN."
                )
                resolved, unresolved = self._resolve_missing_offer_ids(
                    api_url,
                    auth_url,
                    client_id,
                    client_secret,
                    only_successful=True,
                )
                self._log(f"CREATE lookup offer-id: uzupełnione={resolved} | nadal brak={unresolved}")

            try:
                written, missing_after_send = self._persist_offer_ids_to_xlsx(mapping, only_successful=True)
                if written:
                    self._log(f"XLSX: zapisano/uzupełniono offer-id dla {written} wierszy w pliku źródłowym.")
                if missing_after_send:
                    self._log(f"XLSX: nadal brak offer-id dla {missing_after_send} wysłanych wierszy CREATE.")
            except Exception as exc:
                self._log(f"Błąd zapisu offer-id do XLSX: {exc}")

        self._log(f"Koniec wysyłki. Sukces: {ok_count} | Błędy: {fail_count}")
        if fail_count:
            self._set_status(f"Wysyłka zakończona z błędami. Sukces: {ok_count}, błędy: {fail_count}", is_error=True)
        else:
            self._set_status(f"Wysyłka zakończona. Sukces: {ok_count}")


if __name__ == "__main__":
    app = OffersCreatorApp()
    app.mainloop()
